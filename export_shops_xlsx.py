import sys
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_shops_xlsx(input_json_path, output_xlsx_path):
    with open(input_json_path, 'r', encoding='utf-8') as f:
        shops = json.load(f)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Shops Directory"
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10, color="111827")
    mono_font = Font(name="Consolas", size=10, color="111827")
    bold_font = Font(name="Segoe UI", size=10, bold=True, color="111827")
    zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    headers = [
        "S.No",
        "Shop No",
        "Floor / Section",
        "Current Brand",
        "Current Allottee / Owner",
        "Area (Sqft)",
        "Monthly Rent (PKR)",
        "CAM Charges (PKR)",
        "Start Date",
        "Tenure History & Timeline"
    ]

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    for idx, shop in enumerate(shops, 1):
        row_num = idx + 1
        history_str = ""
        if shop.get("history") and len(shop["history"]) > 0:
            hist_parts = []
            for h in shop["history"]:
                p = h.get("period", "")
                b = h.get("brand", "")
                a = h.get("allottee", "")
                if b and a and b != a:
                    hist_parts.append(f"[{p}] {b} (Allottee: {a})")
                elif b:
                    hist_parts.append(f"[{p}] {b}")
                elif a:
                    hist_parts.append(f"[{p}] {a}")
            history_str = " | ".join(hist_parts)

        row_data = [
            idx,
            shop.get("shop_no", ""),
            shop.get("floor", ""),
            shop.get("brand", ""),
            shop.get("allottee", ""),
            shop.get("sqft", ""),
            shop.get("monthly_rent", ""),
            shop.get("cam", ""),
            shop.get("start_date", ""),
            history_str
        ]
        ws.append(row_data)

        is_even = (idx % 2 == 0)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            if is_even:
                cell.fill = zebra_fill

            # Alignment specific rules
            if col_idx == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 2:
                cell.font = bold_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in (6, 7, 8):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_idx == 9:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 10:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        ws.row_dimensions[row_num].height = 22

    # Freeze header row
    ws.freeze_panes = "A2"

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            v = str(cell.value or '')
            if len(v) > max_len:
                max_len = len(v)
        # Cap max_len for history
        if col_letter == 'J':
            ws.column_dimensions[col_letter].width = 45
        else:
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(output_xlsx_path)
    print(f"Excel export generated successfully: {output_xlsx_path}")

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python export_shops_xlsx.py <input_json> <output_xlsx>")
        sys.exit(1)
    generate_shops_xlsx(sys.argv[1], sys.argv[2])
